"""Parser routing by detected format (C-073) into the canonical ParsedDocument (C-086).

Native parsers keep structure: headings -> section tree (C-084), page numbers (C-085),
tables as markdown (C-077), slides + speaker notes, spreadsheet rows with headers (C-079),
HTML without boilerplate (C-080), code split by definitions (C-081), email/zip
attachments recursively (C-083). PDFs and images go to Document Intelligence when
configured (C-074..C-078); audio/video to the transcriber (C-082).
"""

import ast
import asyncio
import csv
import email
import hashlib
import io
import json
import re
import zipfile
from email import policy
from html.parser import HTMLParser

from app.core.exceptions import DocumentLoadError
from app.core.logging import get_logger
from app.rag.ingestion.filetypes import IMAGE_FORMATS, detect_type, image_format
from app.rag.ingestion.language import detect_language
from app.rag.ingestion.models import Element, ParsedDocument, SourceDocument

logger = get_logger(__name__)

MAX_ARCHIVE_DEPTH = 2
MAX_ARCHIVE_MEMBERS = 500


# Upload formats the parser router handles (shown to admins; detection is by content).
SUPPORTED_EXTENSIONS = [
    ".pdf", ".docx", ".pptx", ".xlsx", ".html", ".htm", ".md", ".txt", ".csv", ".tsv",
    ".json", ".jsonl", ".xml", ".yaml", ".yml", ".eml", ".zip",
    ".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tiff", ".heic",
    ".mp3", ".mp4", ".m4a", ".wav", ".ogg", ".flac",
    ".py", ".js", ".ts", ".java", ".cs", ".go", ".sql", ".sh", ".ps1", ".r",
]  # fmt: skip

class _Sections:
    """Tracks the heading breadcrumb while walking a document."""

    def __init__(self) -> None:
        self.path: list[tuple[int, str]] = []

    def heading(self, level: int, text: str) -> list[str]:
        self.path = [(lvl, t) for lvl, t in self.path if lvl < level] + [(level, text)]
        return self.current

    @property
    def current(self) -> list[str]:
        return [t for _, t in self.path]


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


# -- text formats ---------------------------------------------------------------------------
def parse_text(text: str) -> list[Element]:
    return [Element("paragraph", p.strip()) for p in re.split(r"\n\s*\n", text) if p.strip()]


def parse_markdown(text: str) -> list[Element]:
    elements: list[Element] = []
    sections = _Sections()
    block: list[str] = []
    in_code = False

    def flush(kind: str = "paragraph") -> None:
        if block and "".join(block).strip():
            joined = "\n".join(block).strip()
            if kind == "paragraph" and all(
                line.lstrip().startswith("|") for line in block if line.strip()
            ):
                kind = "table"
            elements.append(Element(kind, joined, section_path=sections.current))
        block.clear()

    for line in text.splitlines():
        if line.strip().startswith("```"):
            if in_code:
                flush("code")
            else:
                flush()
            in_code = not in_code
            continue
        if in_code:
            block.append(line)
            continue
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if heading:
            flush()
            level = len(heading.group(1))
            path = sections.heading(level, heading.group(2))
            elements.append(
                Element(
                    "title" if level == 1 and not elements else "heading",
                    heading.group(2),
                    level=level,
                    section_path=path,
                )
            )
        elif re.match(r"^\s*([-*+]|\d+[.)])\s+", line):
            flush()
            elements.append(
                Element(
                    "list_item",
                    re.sub(r"^\s*([-*+]|\d+[.)])\s+", "", line),
                    section_path=sections.current,
                )
            )
        elif not line.strip():
            flush()
        else:
            block.append(line)
    flush("code" if in_code else "paragraph")
    return elements


class _HtmlToElements(HTMLParser):
    _SKIP = {
        "script",
        "style",
        "nav",
        "footer",
        "header",
        "aside",
        "noscript",
        "form",
        "svg",
        "button",
    }
    _BLOCK = {"p", "div", "section", "article", "blockquote", "dd", "dt", "td", "th", "br"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.elements: list[Element] = []
        self.sections = _Sections()
        self.skip_depth = 0
        self.buffer: list[str] = []
        self.heading_level = 0
        self.in_li = False
        self.in_pre = False
        self.table: list[list[str]] | None = None
        self.cell: list[str] | None = None
        self.title = ""
        self.in_title = False

    def _flush(self, kind: str = "paragraph") -> None:
        text = re.sub(r"[ \t]+", " ", "".join(self.buffer)).strip()
        self.buffer.clear()
        if not text:
            return
        if kind == "heading":
            path = self.sections.heading(self.heading_level, text)
            self.elements.append(
                Element("heading", text, level=self.heading_level, section_path=path)
            )
        else:
            self.elements.append(Element(kind, text, section_path=self.sections.current))

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP:
            self.skip_depth += 1
        if self.skip_depth:
            return
        if tag == "title":
            self.in_title = True
        elif re.fullmatch(r"h[1-6]", tag):
            self._flush()
            self.heading_level = int(tag[1])
        elif tag == "li":
            self._flush()
            self.in_li = True
        elif tag == "pre":
            self._flush()
            self.in_pre = True
        elif tag == "table":
            self._flush()
            self.table = []
        elif tag == "tr" and self.table is not None:
            self.table.append([])
        elif tag in ("td", "th") and self.table is not None:
            self.cell = []
        elif tag in self._BLOCK and self.table is None:
            self._flush("list_item" if self.in_li else "paragraph")

    def handle_endtag(self, tag):
        if tag in self._SKIP:
            self.skip_depth = max(0, self.skip_depth - 1)
            return
        if self.skip_depth:
            return
        if tag == "title":
            self.in_title = False
        elif re.fullmatch(r"h[1-6]", tag):
            self._flush("heading")
            self.heading_level = 0
        elif tag == "li":
            self._flush("list_item")
            self.in_li = False
        elif tag == "pre":
            self._flush("code")
            self.in_pre = False
        elif tag in ("td", "th") and self.table is not None and self.cell is not None:
            if self.table:
                self.table[-1].append(" ".join("".join(self.cell).split()).replace("|", "\\|"))
            self.cell = None
        elif tag == "table" and self.table is not None:
            rows = [r for r in self.table if any(c for c in r)]
            if rows:
                width = max(len(r) for r in rows)
                lines = ["| " + " | ".join(r + [""] * (width - len(r))) + " |" for r in rows]
                lines.insert(1, "|" + " --- |" * width)
                self.elements.append(
                    Element("table", "\n".join(lines), section_path=self.sections.current)
                )
            self.table = None
        elif tag in self._BLOCK and self.table is None:
            self._flush("list_item" if self.in_li else "paragraph")

    def handle_data(self, data):
        if self.skip_depth:
            return
        if self.in_title:
            self.title += data
        elif self.cell is not None:
            self.cell.append(data)
        else:
            self.buffer.append(data if self.in_pre else data.replace("\n", " "))

    def close(self):
        super().close()
        self._flush()


def parse_html(text: str) -> tuple[str, list[Element]]:
    parser = _HtmlToElements()
    parser.feed(text)
    parser.close()
    return parser.title.strip(), parser.elements


def parse_delimited(text: str, delimiter: str) -> list[Element]:
    """C-079: each row rendered with its column headers so it stands alone in a chunk."""
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    if not rows:
        return []
    header = [h.strip() for h in rows[0]]
    return [
        Element(
            "table",
            "; ".join(f"{h}: {v.strip()}" for h, v in zip(header, row, strict=False) if v.strip()),
        )
        for row in rows[1:]
        if any(v.strip() for v in row)
    ]


def parse_json(text: str) -> list[Element]:
    try:
        data = json.loads(text)
    except ValueError:
        return [Element("paragraph", line) for line in text.splitlines() if line.strip()]
    items = data if isinstance(data, list) else [data]
    return [Element("paragraph", json.dumps(item, ensure_ascii=False, indent=1)) for item in items]


# -- office formats (C-079) -----------------------------------------------------------------
def parse_docx(content: bytes) -> tuple[str, list[Element]]:
    import docx

    document = docx.Document(io.BytesIO(content))
    elements: list[Element] = []
    sections = _Sections()
    title = document.core_properties.title or ""
    body = document.element.body
    paragraphs = {p._p: p for p in document.paragraphs}
    tables = {t._tbl: t for t in document.tables}
    for child in body.iterchildren():
        if child in paragraphs:
            paragraph = paragraphs[child]
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
            match = re.match(r"heading (\d)", style)
            if style == "title" or match:
                # The document title sits above Heading 1 in the section breadcrumb.
                level = 0 if style == "title" else int(match.group(1))
                path = sections.heading(level, text)
                title = title or (text if level <= 1 else "")
                elements.append(
                    Element(
                        "title" if style == "title" else "heading",
                        text,
                        level=level,
                        section_path=path,
                    )
                )
            elif "list" in style:
                elements.append(Element("list_item", text, section_path=sections.current))
            else:
                elements.append(Element("paragraph", text, section_path=sections.current))
        elif child in tables:
            rows = [
                [c.text.strip().replace("|", "\\|") for c in row.cells]
                for row in tables[child].rows
            ]
            if rows:
                lines = ["| " + " | ".join(r) + " |" for r in rows]
                lines.insert(1, "|" + " --- |" * len(rows[0]))
                elements.append(Element("table", "\n".join(lines), section_path=sections.current))
    return title, elements


def parse_pptx(content: bytes) -> tuple[str, list[Element]]:
    from pptx import Presentation

    deck = Presentation(io.BytesIO(content))
    elements: list[Element] = []
    title = deck.core_properties.title or ""
    for number, slide in enumerate(deck.slides, start=1):
        slide_title = (
            slide.shapes.title.text.strip()
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame
            else ""
        )
        texts = [
            shape.text_frame.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame
            and shape != slide.shapes.title
            and shape.text_frame.text.strip()
        ]
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        body = "\n".join(texts) + (f"\nSpeaker notes: {notes}" if notes else "")
        heading = slide_title or f"Slide {number}"
        title = title or (slide_title if number == 1 else "")
        elements.append(
            Element("slide", f"{heading}\n{body}".strip(), page=number, section_path=[heading])
        )
    return title, elements


def parse_xlsx(content: bytes) -> list[Element]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    elements: list[Element] = []
    for sheet in workbook.worksheets:
        rows = [
            [("" if v is None else str(v)) for v in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        rows = [r for r in rows if any(c.strip() for c in r)]
        if not rows:
            continue
        header = rows[0]
        for row in rows[1:]:
            record = "; ".join(f"{h}: {v}" for h, v in zip(header, row, strict=False) if v.strip())
            if record:
                elements.append(Element("table", record, section_path=[sheet.title]))
    return elements


def parse_pdf_native(content: bytes) -> tuple[list[Element], int, list[int]]:
    """C-074: text per page in reading order with page numbers; returns pages lacking text."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    elements: list[Element] = []
    empty_pages: list[int] = []
    for number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if len(text.strip()) < 30:
            empty_pages.append(number)
        for paragraph in re.split(r"\n\s*\n", text):
            if paragraph.strip():
                elements.append(Element("paragraph", paragraph.strip(), page=number))
    return elements, len(reader.pages), empty_pages


_VISION_MAX_SIDE = 2048  # larger images are downscaled before sending (the model resizes anyway)


def extract_pdf_images(
    content: bytes, min_pixels: int, max_images: int
) -> tuple[list[tuple[int, bytes]], int]:
    """C-078: embedded images worth describing, as (page, PNG or JPEG bytes) in page order.

    Skips images whose shorter side is under `min_pixels` (icons, bullets, logos) and repeats
    of an image already taken (a logo in every page header). Returns the images and how many
    more qualified beyond `max_images`.
    """
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    found: list[tuple[int, bytes]] = []
    seen: set[str] = set()
    over_limit = 0
    for number, page in enumerate(reader.pages, start=1):
        try:
            images = list(page.images)
        except Exception as exc:  # noqa: BLE001 - one unreadable page must not fail the PDF
            logger.warning("pdf_page_images_unreadable", page=number, error=str(exc))
            continue
        for image in images:
            try:
                picture = image.image
                data = image.data
            except Exception as exc:  # noqa: BLE001 - unsupported filters (JBIG2, ...)
                logger.warning("pdf_image_unreadable", page=number, error=str(exc))
                continue
            if picture is None or min(picture.size) < min_pixels:
                continue
            digest = hashlib.sha256(data).hexdigest()
            if digest in seen:
                continue
            seen.add(digest)
            if len(found) >= max_images:
                over_limit += 1
                continue
            if (
                image_format(data) not in ("png", "jpeg")
                or picture.mode not in ("RGB", "L")
                or max(picture.size) > _VISION_MAX_SIDE
            ):
                picture = picture.convert("RGB")
                picture.thumbnail((_VISION_MAX_SIDE, _VISION_MAX_SIDE))
                buffer = io.BytesIO()
                picture.save(buffer, "PNG")
                data = buffer.getvalue()
            found.append((number, data))
    return found, over_limit


# -- code (C-081) ---------------------------------------------------------------------------
def parse_code(text: str, extension: str) -> list[Element]:
    """One element per top-level definition so functions/classes stay whole.
    Python uses its AST; other languages a definition-keyword split (tree-sitter TODO)."""
    if extension == ".py":
        try:
            tree = ast.parse(text)
        except SyntaxError:
            return [Element("code", text)]
        lines = text.splitlines()
        elements = []
        for node in tree.body:
            segment = "\n".join(lines[node.lineno - 1 : node.end_lineno])
            name = getattr(node, "name", "")
            elements.append(Element("code", segment, section_path=[name] if name else []))
        return elements or [Element("code", text)]
    parts = re.split(
        r"\n(?=(?:public |private |export |async "
        r")?(?:function|class|def|func|interface|CREATE|ALTER)\b)",
        text,
    )
    return [Element("code", p.strip()) for p in parts if p.strip()]


# -- router ---------------------------------------------------------------------------------
class ParserRouter:
    def __init__(
        self,
        doc_intelligence=None,
        vision=None,
        transcriber=None,
        pdf_parser: str = "auto",
        image_parser: str = "auto",
        image_min_ocr_chars: int = 30,
        pdf_image_descriptions: bool = True,
        pdf_image_min_pixels: int = 150,
        pdf_max_images: int = 30,
        vision_concurrency: int = 4,
    ) -> None:
        self._docintel = doc_intelligence
        self._vision = vision
        self._transcriber = transcriber
        self._pdf_parser = pdf_parser  # native | docintel | auto (docintel when configured)
        self._image_parser = image_parser  # see Settings.image_parser
        self._image_min_ocr_chars = image_min_ocr_chars
        # Images inside natively parsed PDFs go to the vision model (see _describe_pdf_images).
        self._pdf_images = (
            pdf_image_descriptions and vision is not None and image_parser != "off"
        )
        self._pdf_image_min_pixels = pdf_image_min_pixels
        self._pdf_max_images = pdf_max_images
        self._vision_slots = asyncio.Semaphore(vision_concurrency)

    async def parse(self, document: SourceDocument, depth: int = 0) -> ParsedDocument:
        kind = detect_type(document.content, document.name)
        title = document.metadata.get("title") or document.name.rsplit(".", 1)[0]
        parsed = await self._parse_kind(kind, document, title, depth)
        parsed.language = detect_language(parsed.text)
        score_quality(parsed)
        return parsed

    async def _parse_kind(
        self, kind: str, doc: SourceDocument, title: str, depth: int
    ) -> ParsedDocument:
        source, content = doc.source, doc.content

        def make(elements, parser, new_title=None, pages=None) -> ParsedDocument:
            return ParsedDocument(
                source=source,
                title=new_title or title,
                elements=elements,
                parser=parser,
                page_count=pages,
            )

        if kind == ".pdf":
            if self._docintel is not None and self._pdf_parser in ("auto", "docintel"):
                return await self._docintel.parse(source, doc.name, content)
            elements, pages, empty = await asyncio.to_thread(parse_pdf_native, content)
            parser, flags = "pypdf", []
            if self._pdf_images:
                figures, figure_flags = await self._describe_pdf_images(doc.name, content)
                if figures:
                    # Stable sort: each page's text first, then what its images show.
                    elements = sorted([*elements, *figures], key=lambda e: e.page or 0)
                    parser = "pypdf+vision"
                    described = {f.page for f in figures}
                    empty = [p for p in empty if p not in described]
                flags += figure_flags
            parsed = make(elements, parser, pages=pages)
            if empty:
                flags.append(f"pages_without_text:{len(empty)}/{pages}")  # C-075: needs OCR
            parsed.quality_flags.extend(flags)
            return parsed
        if kind == ".docx":
            doc_title, elements = await asyncio.to_thread(parse_docx, content)
            return make(elements, "python-docx", doc_title)
        if kind == ".pptx":
            deck_title, elements = await asyncio.to_thread(parse_pptx, content)
            return make(elements, "python-pptx", deck_title, pages=len(elements))
        if kind == ".xlsx":
            return make(await asyncio.to_thread(parse_xlsx, content), "openpyxl")
        if kind == ".html":
            html_title, elements = parse_html(_decode(content))
            return make(elements, "html", html_title)
        if kind == ".md":
            elements = parse_markdown(_decode(content))
            first = next((e.text for e in elements if e.type == "title"), None)
            return make(elements, "markdown", first)
        if kind in (".csv", ".tsv"):
            return make(parse_delimited(_decode(content), "\t" if kind == ".tsv" else ","), "csv")
        if kind in (".json", ".jsonl"):
            return make(parse_json(_decode(content)), "json")
        if kind in (".py", ".js", ".ts", ".java", ".cs", ".go", ".sql", ".sh", ".ps1", ".r"):
            return make(parse_code(_decode(content), kind), "code")
        if kind == ".eml":
            return self._parse_email(doc, title, depth)
        if kind == ".zip":
            return self._parse_zip(doc, title, depth)
        if kind == ".image":
            return await self._parse_image(doc, title)
        if kind == ".media":
            if self._transcriber is None:
                raise DocumentLoadError("Audio/video needs AZURE_OPENAI_WHISPER_DEPLOYMENT")
            return await self._transcriber.parse(source, doc.name, content)
        if kind in (".txt", ".xml", ".yaml", ".yml"):
            return make(parse_text(_decode(content)), "text")
        raise DocumentLoadError(
            f"Unsupported file type for {doc.name} (detected {kind or 'unknown'})"
        )

    async def _describe_pdf_images(
        self, name: str, content: bytes
    ) -> tuple[list[Element], list[str]]:
        """C-075 / C-078 without Document Intelligence: what the PDF's diagrams, charts and
        scanned pages show, as one figure element per image on its page. A failed image is
        logged and skipped; the rest of the PDF is still ingested."""
        try:
            images, over_limit = await asyncio.to_thread(
                extract_pdf_images, content, self._pdf_image_min_pixels, self._pdf_max_images
            )
        except Exception as exc:  # noqa: BLE001 - fall back to the text alone
            logger.warning("pdf_image_extraction_failed", document=name, error=str(exc))
            return [], ["pdf_images_unreadable"]

        async def describe(page: int, data: bytes) -> Element | None:
            async with self._vision_slots:
                try:
                    text = (await self._vision.describe(data)).strip()
                except Exception as exc:  # noqa: BLE001 - one image must not fail the PDF
                    logger.warning(
                        "pdf_image_description_failed", document=name, page=page, error=str(exc)
                    )
                    return None
            return Element("figure", f"Figure (page {page}): {text}", page=page) if text else None

        results = await asyncio.gather(*(describe(page, data) for page, data in images))
        figures = [r for r in results if r is not None]
        flags = []
        if over_limit:
            flags.append(f"pdf_images_not_described:{over_limit}")  # raise PDF_MAX_IMAGES
        if len(figures) < len(images):
            flags.append(f"pdf_image_descriptions_failed:{len(images) - len(figures)}")
        logger.info(
            "pdf_images_described",
            document=name,
            described=len(figures),
            extracted=len(images),
            over_limit=over_limit,
        )
        return figures, flags

    async def _parse_image(self, doc: SourceDocument, title: str) -> ParsedDocument:
        """C-075 / C-078: OCR with Document Intelligence and/or describe with the vision model."""
        mode = self._image_parser
        if mode == "off":
            raise DocumentLoadError("Image ingestion is disabled (IMAGE_PARSER=off)")
        fmt = image_format(doc.content)
        _, ocr_supported, vision_supported = IMAGE_FORMATS.get(fmt, ("", False, False))
        can_ocr = self._docintel is not None and ocr_supported
        can_describe = self._vision is not None and vision_supported

        ocr = None
        if can_ocr and mode in ("auto", "docintel", "docintel+vision"):
            ocr = await self._docintel.parse_image(doc.source, doc.name, doc.content)
            enough_text = len(ocr.text.strip()) >= self._image_min_ocr_chars
            if mode == "docintel" or (mode == "auto" and (enough_text or not can_describe)):
                return ocr
        description = None
        if can_describe and mode in ("auto", "vision", "docintel+vision"):
            description = await self._vision.parse(doc.source, doc.name, doc.content)
        if ocr is None and description is None:
            raise DocumentLoadError(
                f"No image processor for {fmt or 'unknown'} image {doc.name} (IMAGE_PARSER={mode})."
                " Document Intelligence reads PNG/JPEG/BMP/TIFF/HEIF (DOC_INTELLIGENCE_ENDPOINT);"
                " the vision model reads PNG/JPEG/GIF/WebP (AZURE_OPENAI_VISION_DEPLOYMENT)."
            )
        if ocr is None:
            return description
        if description is None:
            return ocr
        ocr.elements.extend(description.elements)  # scanned text + what the picture shows
        ocr.parser = f"{ocr.parser}+vision"
        return ocr

    def _parse_email(self, doc: SourceDocument, title: str, depth: int) -> ParsedDocument:
        message = email.message_from_bytes(doc.content, policy=policy.default)
        header = "\n".join(
            f"{h}: {message[h]}" for h in ("From", "To", "Date", "Subject") if message[h]
        )
        body_part = message.get_body(preferencelist=("plain", "html"))
        body = body_part.get_content() if body_part is not None else ""
        if body_part is not None and body_part.get_content_type() == "text/html":
            body = "\n\n".join(e.text for e in parse_html(body)[1])
        elements = [Element("paragraph", header), *parse_text(body)]
        parsed = ParsedDocument(doc.source, str(message["Subject"] or title), elements, "email")
        if depth < MAX_ARCHIVE_DEPTH:
            for part in message.iter_attachments():
                payload = part.get_payload(decode=True) or b""
                name = part.get_filename() or "attachment"
                parsed.attachments.append(
                    SourceDocument(
                        source=f"{doc.source}#{name}",
                        name=name,
                        content=payload,
                        acl=list(doc.acl),
                        url=doc.url,
                    )
                )
        return parsed

    def _parse_zip(self, doc: SourceDocument, title: str, depth: int) -> ParsedDocument:
        parsed = ParsedDocument(doc.source, title, [], "zip")
        if depth >= MAX_ARCHIVE_DEPTH:
            return parsed
        with zipfile.ZipFile(io.BytesIO(doc.content)) as archive:
            members = [m for m in archive.infolist() if not m.is_dir()][:MAX_ARCHIVE_MEMBERS]
            for member in members:
                parsed.attachments.append(
                    SourceDocument(
                        source=f"{doc.source}#{member.filename}",
                        name=member.filename.rsplit("/", 1)[-1],
                        content=archive.read(member),
                        acl=list(doc.acl),
                        url=doc.url,
                    )
                )
        return parsed


def score_quality(parsed: ParsedDocument) -> None:
    """C-088: flag garbled or empty extractions (bad OCR, encoding issues, image-only PDFs)."""
    text = parsed.text
    if not text.strip():
        parsed.quality = 0.0
        parsed.quality_flags.append("no_text")
        return
    letters = sum(c.isalpha() for c in text)
    printable_ratio = sum(c.isprintable() or c in "\n\t" for c in text) / len(text)
    replacement_ratio = text.count("�") / len(text)
    words = re.findall(r"\S+", text)
    avg_word = sum(len(w) for w in words) / len(words) if words else 0
    score = 1.0
    if letters / len(text) < 0.4:
        score -= 0.3
        parsed.quality_flags.append("low_letter_ratio")
    if printable_ratio < 0.95 or replacement_ratio > 0.01:
        score -= 0.4
        parsed.quality_flags.append("garbled_characters")
    if avg_word > 15 or avg_word < 2:
        score -= 0.3
        parsed.quality_flags.append("abnormal_word_length")
    parsed.quality = max(0.0, round(score, 2))
